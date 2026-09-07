"""Cases an adversarial pass found. Each one shipped as a bug once.

Every test here names the wrong behaviour it prevents, so a future change that
reintroduces it fails with an explanation rather than a bare assertion.
"""

from __future__ import annotations

from openpyxl import Workbook

from show_your_work.cli import run


def _book(tmp_path, name: str, build) -> str:
    wb = Workbook()
    build(wb)
    path = tmp_path / name
    wb.save(path)
    return str(path)


def checks(path: str) -> list[str]:
    return [f.check for f in run(path).findings]


def at(path: str, check: str) -> list:
    return [f for f in run(path).findings if f.check == check]


# --- 1. an external reference must not be re-read as a local one ---


def test_external_reference_does_not_invent_a_circular_reference(tmp_path):
    def build(wb):
        first = wb.active
        first.title = "Sheet1"
        second = wb.create_sheet("Sheet2")
        # Stripping only the [Book] part would leave "Sheet2!$B$1", which reads as a
        # local reference and closes a loop that does not exist.
        first["A1"] = "=[OtherBook.xlsx]Sheet2!$B$1"
        second["B1"] = "=Sheet1!A1"

    path = _book(tmp_path, "external.xlsx", build)
    assert "circular_reference" not in checks(path)
    assert "external_link" in checks(path)


def test_a_real_cross_sheet_cycle_is_still_caught(tmp_path):
    def build(wb):
        first = wb.active
        first.title = "A"
        second = wb.create_sheet("B")
        first["A1"] = "=B!B1"
        second["B1"] = "=A!A1"

    assert "circular_reference" in checks(_book(tmp_path, "cycle.xlsx", build))


# --- 2 and 3. a seed value is not an overwritten formula ---


def test_opening_balance_is_not_reported_as_an_overwritten_formula(tmp_path):
    def build(wb):
        ws = wb.active
        ws.title = "Ledger"
        ws["C2"] = 1000  # opening balance, typed on purpose
        for row in range(3, 30):
            ws.cell(row, 2, row * 7)
            ws.cell(row, 3, f"=C{row - 1}+B{row}")

    path = _book(tmp_path, "balance.xlsx", build)
    assert run(path).findings == [], "a running balance is correct spreadsheet practice"


def test_override_in_the_middle_of_a_column_is_still_high(tmp_path):
    def build(wb):
        ws = wb.active
        for row in range(2, 12):
            ws.cell(row, 1, row)
            ws.cell(row, 2, f"=A{row}*2")
        ws["B6"] = 999

    hits = at(_book(tmp_path, "override.xlsx", build), "overwritten_formula")
    assert [f.location for f in hits] == ["B6"]
    assert hits[0].level.value == "high"


def test_override_at_the_edge_is_medium_because_a_typed_edge_is_often_deliberate(tmp_path):
    def build(wb):
        ws = wb.active
        ws["B2"] = 4200  # nothing below refers back to it, so it is not a seed
        for row in range(3, 15):
            ws.cell(row, 1, row)
            ws.cell(row, 2, f"=A{row}*2")

    hits = at(_book(tmp_path, "edge.xlsx", build), "overwritten_formula")
    assert [f.location for f in hits] == ["B2"]
    assert hits[0].level.value == "medium", "README documents the edge case as medium"


# --- 4. stacked section subtotals are correct, not a gap ---


def test_grand_total_over_section_subtotals_is_not_a_gap(tmp_path):
    def build(wb):
        ws = wb.active
        ws.title = "Rollup"
        for row, value in ((1, 10), (2, 20), (3, 30)):
            ws.cell(row, 2, value)
        ws["B4"] = "=SUM(B1:B3)"
        for row, value in ((5, 40), (6, 50), (7, 60)):
            ws.cell(row, 2, value)
        ws["B8"] = "=SUM(B5:B7)"  # correctly excludes the subtotal in B4
        ws["B9"] = "=B4+B8"

    path = _book(tmp_path, "rollup.xlsx", build)
    assert run(path).findings == [], "stacked sections are how a budget is meant to be built"


# --- 5. the horizontal axis is checked too ---


def test_row_sum_that_misses_the_last_column(tmp_path):
    def build(wb):
        ws = wb.active
        ws.title = "Quarters"
        for col, value in ((2, 100), (3, 110), (4, 120), (5, 130)):
            ws.cell(1, col, value)
        ws["F1"] = "=SUM(B1:D1)"  # misses column E

    hits = at(_book(tmp_path, "quarters.xlsx", build), "total_misses_rows")
    assert [f.location for f in hits] == ["F1"]
    assert "column" in hits[0].summary


def test_column_sum_gap_still_found(tmp_path):
    def build(wb):
        ws = wb.active
        for row in range(2, 10):
            ws.cell(row, 1, row * 5)
        ws["A10"] = "=SUM(A2:A7)"

    hits = at(_book(tmp_path, "vertical.xlsx", build), "total_misses_rows")
    assert [f.location for f in hits] == ["A10"]
    assert "2 rows" in hits[0].summary


# --- error reporting ---


def test_zip_without_a_workbook_gets_a_plain_message(tmp_path, capsys):
    import zipfile

    from show_your_work.cli import EXIT_UNREADABLE, main

    path = tmp_path / "empty.xlsx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("readme.txt", "not a workbook")
    assert main([str(path)]) == EXIT_UNREADABLE
    message = capsys.readouterr().err
    assert "does not contain a workbook" in message
    assert "KeyError" not in message
