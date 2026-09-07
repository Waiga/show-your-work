"""Hostile-file refusals: a workbook must not be able to exhaust memory or smuggle
an XML entity-expansion payload past the parser.

Every fixture here is built in code from invented values and repacked in the same
process, so the suite carries no spreadsheet on disk and reproduces each attack from
its parts. Each test names the wrong behaviour it prevents: a traceback, a silent
empty report, or an out-of-memory kill where a plain refusal belongs.
"""

from __future__ import annotations

import zipfile

from openpyxl import Workbook

from show_your_work import workbook as wb_module
from show_your_work.cli import EXIT_CLEAN, EXIT_UNREADABLE, main
from show_your_work.workbook import UnreadableWorkbook, load


def _base(tmp_path, name: str = "base.xlsx") -> str:
    """A small, ordinary workbook to graft a hostile part onto."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Model"
    ws["A1"] = "Label"
    for row in range(2, 6):
        ws.cell(row, 1, "item")
        ws.cell(row, 2, row * 10)
    path = tmp_path / name
    wb.save(path)
    return str(path)


def _repack(src: str, dst, part: str, data: bytes) -> str:
    """Copy every part of ``src`` except ``part``, which is replaced by ``data``."""
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        names = zin.namelist()
        for name in names:
            zout.writestr(name, data if name == part else zin.read(name))
        if part not in names:
            zout.writestr(part, data)
    return str(dst)


# --- (a) XML entity expansion ---------------------------------------------------


def _entity_payload() -> bytes:
    """A sharedStrings part that declares nested entities (the billion-laughs shape)."""
    decls = ['<!ENTITY a0 "{}">'.format("A" * 100)]
    for i in range(1, 10):
        decls.append('<!ENTITY a{} "{}">'.format(i, f"&a{i - 1};" * 10))
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        "<!DOCTYPE sst [\n" + "\n".join(decls) + "\n]>\n"
        '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'count="1" uniqueCount="1"><si><t>&a9;</t></si></sst>'
    ).encode()


def test_entity_declaration_is_refused_unread(tmp_path):
    path = _repack(
        _base(tmp_path), tmp_path / "entity.xlsx", "xl/sharedStrings.xml", _entity_payload()
    )
    try:
        load(path)
    except UnreadableWorkbook as exc:
        assert "entity" in str(exc).lower()
    else:
        raise AssertionError("a workbook declaring XML entities must be refused")


def test_entity_workbook_exits_two_with_a_plain_message(tmp_path, capsys):
    path = _repack(
        _base(tmp_path), tmp_path / "entity.xlsx", "xl/sharedStrings.xml", _entity_payload()
    )
    assert main([path]) == EXIT_UNREADABLE
    err = capsys.readouterr().err
    assert "entity" in err.lower()
    # never a traceback, never a stray exception name
    assert "Traceback" not in err
    assert "ParseError" not in err


def test_a_lone_doctype_without_entities_is_still_refused(tmp_path):
    """A DOCTYPE alone is the machinery an attack needs and no workbook writes it."""
    payload = (
        b'<?xml version="1.0"?>\n<!DOCTYPE worksheet>\n'
        b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b"<sheetData/></worksheet>"
    )
    path = _repack(_base(tmp_path), tmp_path / "doctype.xlsx", "xl/worksheets/sheet1.xml", payload)
    try:
        load(path)
    except UnreadableWorkbook as exc:
        assert "document type" in str(exc).lower() or "entity" in str(exc).lower()
    else:
        raise AssertionError("a DOCTYPE declaration must be refused")


# --- (b) zip decompression bomb -------------------------------------------------


def _oversized_sheet(uncompressed_bytes: int) -> bytes:
    """A valid sheet XML padded to a chosen size, so it is on the parse path."""
    return (
        b'<?xml version="1.0"?>'
        b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b"<!--" + b"\x20" * uncompressed_bytes + b"-->"
        b"<sheetData/></worksheet>"
    )


def test_a_single_hugely_compressible_part_is_refused(tmp_path):
    # 200 MB of spaces compresses to a few hundred KB: ratio well over the limit.
    payload = _oversized_sheet(200 * 1024 * 1024)
    path = _repack(_base(tmp_path), tmp_path / "bomb.xlsx", "xl/worksheets/sheet1.xml", payload)
    try:
        load(path)
    except UnreadableWorkbook as exc:
        assert "decompression bomb" in str(exc)
    else:
        raise AssertionError("a high-ratio part must be refused before it is parsed")


def test_zip_bomb_exits_two_and_does_not_read_the_file(tmp_path, capsys):
    payload = _oversized_sheet(200 * 1024 * 1024)
    path = _repack(_base(tmp_path), tmp_path / "bomb.xlsx", "xl/worksheets/sheet1.xml", payload)
    assert main([path]) == EXIT_UNREADABLE
    err = capsys.readouterr().err
    # A refusal must read as "nothing was examined", never as an empty clean report.
    assert "refused unread" in err
    assert "Traceback" not in err


def test_total_uncompressed_size_over_the_limit_is_refused(tmp_path, monkeypatch):
    """The total-size guard, exercised with a low limit so the test stays fast."""
    monkeypatch.setattr(wb_module, "MAX_TOTAL_UNCOMPRESSED", 4096)
    try:
        load(_base(tmp_path))
    except UnreadableWorkbook as exc:
        assert "unpacks to" in str(exc)
    else:
        raise AssertionError("a workbook over the total-size limit must be refused")


# --- false-positive guards ------------------------------------------------------


def test_an_ordinary_large_workbook_still_opens(tmp_path, capsys):
    """A dense, legitimate workbook must survive both the size and ratio checks."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Model"
    for row in range(1, 4001):
        for col in range(1, 6):
            ws.cell(row, col, row * col)
        ws.cell(row, 6, f"=SUM(A{row}:E{row})")
    path = tmp_path / "big.xlsx"
    wb.save(path)
    assert main([str(path)]) == EXIT_CLEAN


def test_a_repetitive_but_legitimate_workbook_still_opens(tmp_path):
    """Identical strings compress well but stay far under the ratio limit."""
    wb = Workbook()
    ws = wb.active
    for row in range(1, 5001):
        ws.cell(row, 1, "the same label repeated many times over")
        ws.cell(row, 2, 12345)
    path = tmp_path / "repetitive.xlsx"
    wb.save(path)
    # No exception, and it is not mistaken for a bomb.
    loaded = load(str(path))
    loaded.close()


def test_a_small_part_that_compresses_extremely_well_is_not_a_bomb(tmp_path):
    """Ratio alone is not evidence. A part too small to exhaust anything is not judged.

    Ordinary workbooks carry small, highly repetitive parts — a theme, a styles table,
    a sheet of one repeated value — and those can sit well past the ratio limit for
    entirely innocent reasons. Refusing them would cost real files for no protection,
    because many small parts are already caught by the total-size limit.
    """
    base = _base(tmp_path, "innocent.xlsx")
    # About 2 MB of one repeated byte: a ratio in the thousands, and far too small
    # to exhaust anything.
    grafted = _repack(
        base, tmp_path / "compressible.xlsx", "xl/harmless.xml",
        b"<a>" + b"x" * (2 * 1024 * 1024) + b"</a>",
    )
    assert load(grafted).formulas is not None
