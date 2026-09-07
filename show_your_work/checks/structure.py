"""Checks about the shape of the workbook rather than its formulas."""

from __future__ import annotations

from openpyxl.utils import get_column_letter

from show_your_work.findings import Finding, Level, Report
from show_your_work.workbook import LoadedWorkbook

# Runs of hidden rows are reported as ranges; a lone hidden row is usually deliberate,
# a long run of them usually is not.
MANY_HIDDEN = 3


def _group_runs(numbers: list[int]) -> list[tuple[int, int]]:
    """Collapse [2,3,4,9] into [(2,4),(9,9)]."""
    runs: list[tuple[int, int]] = []
    for n in sorted(set(numbers)):
        if runs and n == runs[-1][1] + 1:
            runs[-1] = (runs[-1][0], n)
        else:
            runs.append((n, n))
    return runs


def hidden_sheets(book: LoadedWorkbook, report: Report) -> None:
    """Sheets that do not appear in the tab bar."""
    for sheet in book.formulas.worksheets:
        state = sheet.sheet_state
        if state == "visible":
            continue
        very = state == "veryHidden"
        report.add(
            Finding(
                check="hidden_sheet",
                level=Level.HIGH if very else Level.MEDIUM,
                sheet=sheet.title,
                location="(whole sheet)",
                summary=(
                    "Sheet is hidden and cannot be unhidden from the Excel menu"
                    if very
                    else "Sheet is hidden"
                ),
                detail=(
                    "A 'very hidden' sheet is only reachable through the VBA editor, so a "
                    "reader reviewing this file in Excel has no way to know it is there."
                    if very
                    else "Visible totals may depend on rows a reader never sees."
                ),
            )
        )


def hidden_rows_and_columns(book: LoadedWorkbook, report: Report) -> None:
    """Rows and columns hidden from view inside an otherwise visible sheet."""
    for sheet in book.formulas.worksheets:
        rows = [idx for idx, dim in sheet.row_dimensions.items() if dim.hidden]
        cols = [key for key, dim in sheet.column_dimensions.items() if dim.hidden]

        for start, end in _group_runs(rows):
            span = end - start + 1
            report.add(
                Finding(
                    check="hidden_rows",
                    level=Level.MEDIUM if span >= MANY_HIDDEN else Level.LOW,
                    sheet=sheet.title,
                    location=f"rows {start}-{end}" if span > 1 else f"row {start}",
                    summary=f"{span} hidden row{'s' if span > 1 else ''}",
                    detail="Hidden rows still count toward totals that cover them.",
                )
            )

        if cols:
            report.add(
                Finding(
                    check="hidden_columns",
                    level=Level.MEDIUM if len(cols) >= MANY_HIDDEN else Level.LOW,
                    sheet=sheet.title,
                    location=", ".join(sorted(cols)),
                    summary=f"{len(cols)} hidden column{'s' if len(cols) > 1 else ''}",
                    detail="Hidden columns still count toward totals that cover them.",
                )
            )


def active_filters(book: LoadedWorkbook, report: Report) -> None:
    """A filter means the rows on screen may not be all the rows."""
    for sheet in book.formulas.worksheets:
        ref = getattr(sheet.auto_filter, "ref", None)
        if not ref:
            continue
        has_criteria = bool(getattr(sheet.auto_filter, "filterColumn", None))
        report.add(
            Finding(
                check="active_filter",
                level=Level.MEDIUM if has_criteria else Level.LOW,
                sheet=sheet.title,
                location=str(ref),
                summary=(
                    "A filter with active criteria is applied"
                    if has_criteria
                    else "A filter is set up over this range"
                ),
                detail=(
                    "Rows excluded by the filter are hidden, so a figure read off the "
                    "screen can differ from the figure a formula computes."
                ),
            )
        )


def merged_cells_in_data(book: LoadedWorkbook, report: Report) -> None:
    """Merged cells inside a data block break sorting, filtering and range formulas."""
    for sheet in book.formulas.worksheets:
        merges = list(sheet.merged_cells.ranges)
        if not merges:
            continue
        # A merge spanning several rows sits inside data; a single-row merge is
        # nearly always a title or a header spanning columns.
        multi_row = [m for m in merges if m.max_row > m.min_row]
        if not multi_row:
            continue
        report.add(
            Finding(
                check="merged_cells",
                level=Level.LOW,
                sheet=sheet.title,
                location=", ".join(str(m) for m in multi_row[:5])
                + ("" if len(multi_row) <= 5 else f" (+{len(multi_row) - 5} more)"),
                summary=f"{len(multi_row)} merged block{'s' if len(multi_row) > 1 else ''} "
                "spanning multiple rows",
                detail=(
                    "Only the top-left cell of a merge holds a value, so ranges over "
                    "these rows read blanks where a reader sees a number."
                ),
            )
        )


def macro_enabled_file(book: LoadedWorkbook, report: Report) -> None:
    """Macros can change numbers on open, before anyone reads them."""
    if not book.macro_enabled:
        return
    report.add(
        Finding(
            check="macro_enabled",
            level=Level.MEDIUM,
            sheet="(workbook)",
            location=book.path.name,
            summary="Macro-enabled workbook (.xlsm)",
            detail=(
                "Code stored in this file may alter values when it is opened. "
                "This tool reads the sheet contents only and does not inspect that code."
            ),
        )
    )


def protected_sheets(book: LoadedWorkbook, report: Report) -> None:
    """Sheet protection is recorded, not treated as a problem on its own."""
    for sheet in book.formulas.worksheets:
        protection = getattr(sheet, "protection", None)
        if protection is None or not getattr(protection, "sheet", False):
            continue
        report.add(
            Finding(
                check="protected_sheet",
                level=Level.LOW,
                sheet=sheet.title,
                location="(whole sheet)",
                summary="Sheet is protected",
                detail=(
                    "Protection limits editing in Excel. It does not hide the contents "
                    "from this tool, and every other check still ran on this sheet."
                ),
            )
        )


def _column_letter(index: int) -> str:
    return get_column_letter(index)
