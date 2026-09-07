"""Checks that read the formulas themselves.

Every finding here is anchored to something present in the file. The tool never
decides whether a number is *right*; it reports only that a number cannot be
explained by the sheet it lives in.
"""

from __future__ import annotations

import re

from openpyxl.utils import column_index_from_string, get_column_letter

from show_your_work.findings import Finding, Level, Report
from show_your_work.workbook import LoadedWorkbook

# A run needs enough formulas around a constant before the constant means anything.
MIN_RUN_FOR_CONTEXT = 3
# Guard against a sheet whose declared dimensions are far larger than its content.
MAX_CELLS_PER_SHEET = 400_000

EXCEL_ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#SPILL!")

VOLATILE = ("NOW", "TODAY", "RAND", "RANDBETWEEN", "RANDARRAY", "OFFSET", "INDIRECT")
_VOLATILE_RE = re.compile(r"\b(" + "|".join(VOLATILE) + r")\s*\(", re.IGNORECASE)
_STRING_LITERAL_RE = re.compile(r'"[^"]*"')

# =SUM(C2:C10) and friends, capturing the first plain range argument.
_AGGREGATE_RE = re.compile(
    r"\b(SUM|AVERAGE|COUNT|COUNTA|MIN|MAX|MEDIAN|PRODUCT)\s*\(\s*"
    r"\$?([A-Z]{1,3})\$?(\d+)\s*:\s*\$?([A-Z]{1,3})\$?(\d+)\s*\)",
    re.IGNORECASE,
)

# A subtotal legitimately differs from the column it closes, so it is not an oddity.
# The call is looked for anywhere in the formula, not only at the start: "=+SUM(C2:C9)"
# is the Lotus carry-over idiom and is everywhere in real finance workbooks, and
# "=ROUND(SUM(C2:C9),0)" is the ordinary way to write a rounded total. Requiring the
# aggregate to open the formula reported both of those as broken.
_SUBTOTAL_RE = re.compile(
    r"\b(SUM|SUBTOTAL|AVERAGE|COUNT|COUNTA|MIN|MAX|MEDIAN|PRODUCT|SUMPRODUCT)\s*\(",
    re.IGNORECASE,
)


def _calls_an_aggregate(formula: str) -> bool:
    """True when a formula calls an aggregate anywhere, ignoring text inside quotes.

    The name must be followed directly by its opening bracket, so the conditional
    family — SUMIF, SUMIFS, COUNTIF, COUNTIFS, AVERAGEIF — does not match. Those
    aggregate a filtered subset rather than the block a total closes.
    """
    return bool(_SUBTOTAL_RE.search(_STRING_LITERAL_RE.sub('""', formula)))


def is_subtotal_like(cell, block: list) -> bool:
    """True for an aggregate sitting at either end of the block it summarises.

    Position still decides. Widening the pattern above lets a total be written
    "=+SUM(...)" or "=ROUND(SUM(...),0)", but a cell holding an aggregate part-way
    down a block is not exempt from anything and is still reported.
    """
    if cell.data_type != "f" or not isinstance(cell.value, str):
        return False
    if not _calls_an_aggregate(cell.value):
        return False
    return cell.coordinate in (block[0].coordinate, block[-1].coordinate)


def dominant_shape(formula_cells: list) -> tuple[str | None, list]:
    """The formula shape most cells in a block share, and the cells sharing it."""
    if not formula_cells:
        return None, []
    shapes: dict[str, list] = {}
    for cell in formula_cells:
        shapes.setdefault(_shape(str(cell.value)), []).append(cell)
    return max(shapes.items(), key=lambda kv: len(kv[1]))


# [Book.xlsx]Sheet!A1 or [1]Sheet!A1 both mean "this number comes from another file".
_EXTERNAL_RE = re.compile(r"\[[^\]\[]+\]")

_REF_DIGITS_RE = re.compile(r"(\$?[A-Z]{1,3}\$?)(\d+)")


def _shape(formula: str) -> str:
    """A formula stripped of row numbers, so a filled-down column collapses to one shape."""
    return _REF_DIGITS_RE.sub(r"\1#", formula.upper())


def _is_number_like(cell) -> bool:
    """A cell that participates in arithmetic: a typed number or a formula."""
    if cell.value is None:
        return False
    if cell.data_type == "f":
        return True
    return isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)


def _is_typed_constant(cell) -> bool:
    return cell.value is not None and cell.data_type != "f" and _is_number_like(cell)


def _bounded(sheet) -> tuple[int, int, bool]:
    """Return (max_row, max_col, truncated) with a ceiling on total cells."""
    max_row, max_col = sheet.max_row or 0, sheet.max_column or 0
    if max_row * max_col <= MAX_CELLS_PER_SHEET or max_col == 0:
        return max_row, max_col, False
    return max(1, MAX_CELLS_PER_SHEET // max_col), max_col, True


def _iter_lines(sheet, max_row: int, max_col: int):
    """Yield ('column'|'row', label, [cells]) for every column and every row."""
    for col in range(1, max_col + 1):
        cells = [sheet.cell(row, col) for row in range(1, max_row + 1)]
        yield "column", get_column_letter(col), cells
    for row in range(1, max_row + 1):
        cells = [sheet.cell(row, col) for col in range(1, max_col + 1)]
        yield "row", str(row), cells


def _contiguous_blocks(cells: list) -> list[list]:
    """Split a line into runs of consecutive number-like cells."""
    blocks, current = [], []
    for cell in cells:
        if _is_number_like(cell):
            current.append(cell)
        else:
            if current:
                blocks.append(current)
            current = []
    if current:
        blocks.append(current)
    return blocks



def _references(formula: str, coordinate: str) -> bool:
    """True when a formula names this cell, ignoring $ and case."""
    if not isinstance(formula, str):
        return False
    target = coordinate.replace("$", "").upper()
    body = _STRING_LITERAL_RE.sub('""', formula).replace("$", "").upper()
    for match in re.finditer(r"(?<![A-Z0-9_])([A-Z]{1,3}\d+)(?![A-Z0-9_(])", body):
        if match.group(1) == target:
            return True
    return False


def is_seed_value(cell, block: list, index: dict) -> bool:
    """True for a starting value that the run of formulas is built on.

    A running balance, a cumulative counter and any recursive series all begin with a
    typed number, and the formula next to it refers back to that number. That is the
    sheet working correctly, not a formula someone overwrote.
    """
    position = index[cell.coordinate]
    if position == 0:
        neighbour = block[1] if len(block) > 1 else None
    elif position == len(block) - 1:
        neighbour = block[position - 1]
    else:
        return False
    if neighbour is None or neighbour.data_type != "f":
        return False
    return _references(str(neighbour.value), cell.coordinate)


def overwritten_formulas(book: LoadedWorkbook, report: Report) -> None:
    """A typed number sitting inside a run of formulas.

    This is the finding the tool exists for. Someone replaced a calculation with a
    number, so the sheet no longer explains where that number came from, and it will
    not move when its inputs do.
    """
    seen: dict[tuple[str, str], Level] = {}

    for sheet in book.formulas.worksheets:
        max_row, max_col, truncated = _bounded(sheet)
        if truncated:
            report.note_unchecked(
                f"{sheet.title}: rows beyond {max_row}",
                "The sheet is larger than the scan ceiling, so later rows were not read.",
            )
        if max_row == 0 or max_col == 0:
            continue

        for axis, _label, cells in _iter_lines(sheet, max_row, max_col):
            for block in _contiguous_blocks(cells):
                formulas_in_block = [c for c in block if c.data_type == "f"]
                constants = [c for c in block if _is_typed_constant(c)]
                if len(formulas_in_block) < MIN_RUN_FOR_CONTEXT or not constants:
                    continue
                if len(constants) >= len(formulas_in_block):
                    continue  # mostly typed data with a few formulas: normal input block

                # A subtotal closing the block is not part of the repeated pattern.
                pattern_pool = [c for c in formulas_in_block if not is_subtotal_like(c, block)]
                _, pattern_cells = dominant_shape(pattern_pool)
                if len(pattern_cells) < MIN_RUN_FOR_CONTEXT:
                    continue  # no repeated formula here, so there is no pattern to break

                pattern = {c.coordinate for c in pattern_cells}
                index = {c.coordinate: pos for pos, c in enumerate(block)}
                for cell in constants:
                    if is_seed_value(cell, block, index):
                        continue  # an opening value the formulas below it build on
                    pos = index[cell.coordinate]
                    before = block[pos - 1] if pos > 0 else None
                    after = block[pos + 1] if pos + 1 < len(block) else None
                    surrounded = (
                        before is not None
                        and after is not None
                        and before.coordinate in pattern
                        and after.coordinate in pattern
                    )
                    level = Level.HIGH if surrounded else Level.MEDIUM
                    key = (sheet.title, cell.coordinate)
                    if key in seen and seen[key].rank >= level.rank:
                        continue
                    seen[key] = level
                    report.add(
                        Finding(
                            check="overwritten_formula",
                            level=level,
                            sheet=sheet.title,
                            location=cell.coordinate,
                            summary="Typed number inside a column of formulas"
                            if axis == "column"
                            else "Typed number inside a row of formulas",
                            detail=(
                                f"{len(pattern_cells)} cells in this {axis} share one "
                                "formula and this one does not. The sheet cannot show where "
                                "the number came from, and it will not update when its "
                                "inputs change."
                            ),
                            sample=str(cell.value),
                        )
                    )


def _is_aggregate_cell(cell) -> bool:
    """True for a cell holding SUM, SUBTOTAL and friends.

    Unlike is_subtotal_like this carries no position test, because a subtotal in a
    gap is defined by sitting in the gap. The widened pattern is right here too: a
    section subtotal written "=+SUM(B5:B7)" is the same subtotal as "=SUM(B5:B7)"
    and the grand total above it excludes it for the same reason.
    """
    return (
        cell.data_type == "f"
        and isinstance(cell.value, str)
        and _calls_an_aggregate(cell.value)
    )


def _gap(sheet, fixed: int, first: int, last: int, vertical: bool, formula: str) -> list[int]:
    """Positions in a line that hold their own numbers, skipping other totals.

    A subtotal sitting in the gap is deliberately excluded from the range above it,
    which is how stacked sections in a budget are meant to be built.

    A cell the formula names itself is skipped too. "=SUM(D93:D103)-D104" adds row 104
    outside the range on purpose, which is how a deduction or an adjustment line is
    written; reporting it as a skipped row is reading the formula only as far as its
    first bracket.
    """
    if first < 1 or last < first:
        return []
    found = []
    for pos in range(first, last + 1):
        cell = sheet.cell(pos, fixed) if vertical else sheet.cell(fixed, pos)
        if not _is_number_like(cell) or _is_aggregate_cell(cell):
            continue
        if _references(formula, cell.coordinate):
            continue  # named elsewhere in the same formula, so it is not left out
        found.append(pos)
    return found


# A one-off partial aggregate is a slip; the same one repeated is a design decision.
MIN_REPEATS_FOR_DESIGN = 3


def _shape_counts(sheet, max_row: int, max_col: int) -> dict[str, int]:
    """How many cells on the sheet share each normalised formula shape."""
    counts: dict[str, int] = {}
    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            cell = sheet.cell(row, col)
            if cell.data_type != "f" or not isinstance(cell.value, str):
                continue
            shape = _shape(cell.value)
            counts[shape] = counts.get(shape, 0) + 1
    return counts


def totals_that_miss_rows(book: LoadedWorkbook, report: Report) -> None:
    """An aggregate whose range stops short of the numbers it sits next to.

    Handles a total under a column and a total beside a row. A rectangular range is
    left alone, because "the block it should cover" is not well defined there.
    """
    for sheet in book.formulas.worksheets:
        max_row, max_col, _ = _bounded(sheet)
        if max_row == 0 or max_col == 0:
            continue
        repeats = _shape_counts(sheet, max_row, max_col)

        for row in range(1, max_row + 1):
            for col in range(1, max_col + 1):
                cell = sheet.cell(row, col)
                if cell.data_type != "f" or not isinstance(cell.value, str):
                    continue
                match = _AGGREGATE_RE.search(cell.value)
                if not match:
                    continue
                if repeats.get(_shape(cell.value), 0) >= MIN_REPEATS_FOR_DESIGN:
                    # A summary column that deliberately re-aggregates part of a table
                    # repeats one formula down or across the sheet. A total that really
                    # misses rows is a one-off slip, so repetition is the signal that
                    # the shorter range was chosen rather than forgotten.
                    continue

                func, col1, row1, col2, row2 = match.groups()
                start_row, end_row = sorted((int(row1), int(row2)))
                start_col, end_col = sorted(
                    (column_index_from_string(col1.upper()), column_index_from_string(col2.upper()))
                )

                if start_col == end_col:
                    vertical, fixed = True, start_col
                    first, last, here = start_row, end_row, row
                    unit, line = "row", "column"
                elif start_row == end_row:
                    vertical, fixed = False, start_row
                    first, last, here = start_col, end_col, col
                    unit, line = "column", "row"
                else:
                    continue  # a rectangular range has no single line to compare against

                if here <= last:
                    continue  # the total sits inside or before its own range

                missed = _gap(sheet, fixed, last + 1, here - 1, vertical, cell.value)
                if missed:
                    names = _labels(missed, vertical)
                    report.add(
                        Finding(
                            check="total_misses_rows",
                            level=Level.HIGH,
                            sheet=sheet.title,
                            location=cell.coordinate,
                            summary=f"{func.upper()} skips {len(missed)} {unit}"
                            f"{'s' if len(missed) > 1 else ''} that "
                            f"{'sit' if len(missed) > 1 else 'sits'} between its range "
                            "and the total",
                            detail=(
                                f"The range ends at {unit} {_label(last, vertical)} but "
                                f"{unit}s {names[0]}-{names[-1]} in the same {line} also hold "
                                f"numbers. {unit.capitalize()}s added to a table fall outside "
                                "a total that was never extended."
                            ),
                            sample=cell.value,
                        )
                    )
                    continue

                before = _gap(sheet, fixed, first - 1, first - 1, vertical, cell.value)
                if before:
                    report.add(
                        Finding(
                            check="total_misses_rows",
                            level=Level.MEDIUM,
                            sheet=sheet.title,
                            location=cell.coordinate,
                            summary=f"{func.upper()} may start one {unit} too late",
                            detail=(
                                f"The range starts at {unit} {_label(first, vertical)}, and "
                                f"{unit} {_label(first - 1, vertical)} in the same {line} "
                                "also holds a number that is not itself a total."
                            ),
                            sample=cell.value,
                        )
                    )


def _label(position: int, vertical: bool) -> str:
    return str(position) if vertical else get_column_letter(position)


def _labels(positions: list[int], vertical: bool) -> list[str]:
    return [_label(p, vertical) for p in positions]


def cached_error_values(book: LoadedWorkbook, report: Report) -> None:
    """Formulas whose last saved result was an Excel error."""
    if not book.has_cached_values or book.values is None:
        report.note_unchecked(
            "Formula results (#REF!, #DIV/0! and similar)",
            "This file has no saved results, which happens when it was generated by a "
            "program and never opened in Excel. Open it in Excel, save, and run again to "
            "check the calculated values.",
        )
        return

    for sheet in book.values.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                text = cell.value.strip() if isinstance(cell.value, str) else ""
                # data_type 'e' is openpyxl's own error marker; the text test also catches
                # an error that was stored as plain text.
                if cell.data_type != "e" and text not in EXCEL_ERRORS:
                    continue
                error = text or "an Excel error"
                report.add(
                    Finding(
                        check="error_value",
                        level=Level.HIGH,
                        sheet=sheet.title,
                        location=cell.coordinate,
                        summary=f"Cell holds the Excel error {error}",
                        detail=(
                            "Any total that includes this cell is wrong or is itself an "
                            "error. #REF! in particular means a referenced cell was deleted."
                        ),
                        sample=error,
                    )
                )


def external_workbook_links(book: LoadedWorkbook, report: Report) -> None:
    """Formulas that pull numbers out of another file."""
    for sheet in book.formulas.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type != "f" or not isinstance(cell.value, str):
                    continue
                if not _EXTERNAL_RE.search(cell.value):
                    continue
                report.add(
                    Finding(
                        check="external_link",
                        level=Level.MEDIUM,
                        sheet=sheet.title,
                        location=cell.coordinate,
                        summary="Formula reads from another workbook",
                        detail=(
                            "The number shown is whatever was last saved from the other file. "
                            "If that file moved, was renamed, or changed, this figure is stale "
                            "and the sheet gives no sign of it."
                        ),
                        sample=cell.value,
                    )
                )


def volatile_functions(book: LoadedWorkbook, report: Report) -> None:
    """Functions whose answer changes on their own."""
    for sheet in book.formulas.worksheets:
        hits: list[str] = []
        names: set[str] = set()
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type != "f" or not isinstance(cell.value, str):
                    continue
                found = _VOLATILE_RE.findall(_STRING_LITERAL_RE.sub('""', cell.value))
                if found:
                    hits.append(cell.coordinate)
                    names.update(name.upper() for name in found)
        if not hits:
            continue
        report.add(
            Finding(
                check="volatile_function",
                level=Level.LOW,
                sheet=sheet.title,
                location=", ".join(hits[:5])
                + ("" if len(hits) <= 5 else f" (+{len(hits) - 5} more)"),
                summary=f"{len(hits)} formula{'s' if len(hits) > 1 else ''} using "
                + ", ".join(sorted(names)),
                detail=(
                    "These recalculate on their own, so the file can show different numbers "
                    "tomorrow with no edit in between. A figure quoted from this sheet cannot "
                    "be reproduced later."
                ),
            )
        )


def broken_defined_names(book: LoadedWorkbook, report: Report) -> None:
    """Named ranges pointing at deleted cells."""
    try:
        items = list(book.formulas.defined_names.items())
    except Exception:
        report.note_unchecked("Named ranges", "This workbook's defined names could not be read.")
        return

    for name, defn in items:
        target = str(
            getattr(defn, "attr_text", None) or getattr(defn, "value", "") or ""
        )
        if "#REF" in target.upper():
            report.add(
                Finding(
                    check="broken_defined_name",
                    level=Level.HIGH,
                    sheet="(workbook)",
                    location=str(name),
                    summary=f"Named range '{name}' points at deleted cells",
                    detail=(
                        "Every formula using this name resolves to an error, so any total "
                        "built on it is unusable."
                    ),
                    sample=target,
                )
            )
