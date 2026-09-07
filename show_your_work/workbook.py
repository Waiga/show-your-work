"""Opening a workbook safely, and holding the two views a check may need.

openpyxl can show a cell's formula or its last cached result, never both in one
load. Checks need both, so the file is opened twice and the two views are kept
side by side.
"""

from __future__ import annotations

import contextlib
import zipfile
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook


class UnreadableWorkbook(Exception):
    """The file could not be opened as a modern Excel workbook."""


# Only formats openpyxl actually parses. .xls is a different, older binary format.
READABLE_SUFFIXES = {".xlsx", ".xlsm"}


@dataclass
class LoadedWorkbook:
    """Both views of one file, plus what loading it revealed."""

    path: Path
    formulas: Workbook
    values: Workbook | None
    has_cached_values: bool
    macro_enabled: bool

    @property
    def sheet_titles(self) -> list[str]:
        return list(self.formulas.sheetnames)

    def close(self) -> None:
        for wb in (self.formulas, self.values):
            if wb is not None:
                # Closing must never be the thing that ends a run.
                with contextlib.suppress(Exception):
                    wb.close()


def _describe_open_failure(path: Path, exc: Exception) -> str:
    if path.suffix.lower() == ".xls":
        return (
            "This is the older .xls format, which this tool cannot read. "
            "Re-save it as .xlsx and run again."
        )
    if isinstance(exc, zipfile.BadZipFile):
        return (
            "The file is not a valid .xlsx container. It is either corrupt, mislabelled, "
            "or protected with an open password. These look identical from the outside, "
            "so this tool cannot tell you which."
        )
    if isinstance(exc, KeyError):
        return (
            "The file is a zip archive but does not contain a workbook. It is most "
            "likely corrupt or was renamed from another format."
        )
    return f"The file could not be opened as a workbook: {exc}"


def load(path_str: str) -> LoadedWorkbook:
    """Open a workbook for auditing, or raise UnreadableWorkbook with a plain reason."""
    path = Path(path_str)

    if not path.exists():
        raise UnreadableWorkbook(f"No such file: {path}")
    if path.is_dir():
        raise UnreadableWorkbook(f"{path} is a directory, not a workbook.")
    if path.suffix.lower() == ".xls":
        raise UnreadableWorkbook(
            "This is the older .xls format, which this tool cannot read. "
            "Re-save it as .xlsx and run again."
        )
    if path.suffix.lower() not in READABLE_SUFFIXES:
        raise UnreadableWorkbook(
            f"{path.suffix or 'This file'} is not a format this tool reads. "
            f"Supported: {', '.join(sorted(READABLE_SUFFIXES))}."
        )

    try:
        formulas = load_workbook(path, data_only=False, keep_vba=False)
    except Exception as exc:  # openpyxl raises a wide range of parse errors
        raise UnreadableWorkbook(_describe_open_failure(path, exc)) from exc

    values: Workbook | None
    try:
        values = load_workbook(path, data_only=True, keep_vba=False)
    except Exception:
        # The formula view already opened, so a failure here costs a check, not the run.
        values = None

    has_cached = _detect_cached_values(formulas, values)

    return LoadedWorkbook(
        path=path,
        formulas=formulas,
        values=values,
        has_cached_values=has_cached,
        macro_enabled=_contains_macros(path),
    )


def _contains_macros(path: Path) -> bool:
    """Look for a VBA part inside the container rather than trusting the extension.

    A macro workbook renamed to .xlsx still carries its code, so the file name is not
    evidence either way.
    """
    try:
        with zipfile.ZipFile(path) as archive:
            return any("vbaproject" in name.lower() for name in archive.namelist())
    except Exception:
        return False


def _detect_cached_values(formulas: Workbook, values: Workbook | None) -> bool:
    """True when Excel has saved results alongside the formulas.

    A workbook written by a library and never opened in Excel carries formulas with
    no stored results. Every value-based check is blind on such a file, and saying so
    is more useful than reporting a clean pass.
    """
    if values is None:
        return False

    for sheet in formulas.worksheets:
        if sheet.title not in values.sheetnames:
            continue
        value_sheet = values[sheet.title]
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type != "f":
                    continue
                cached = value_sheet.cell(cell.row, cell.column).value
                if cached is not None:
                    return True
        # Keep scanning other sheets; a single formula sheet may legitimately be empty.
    return False
