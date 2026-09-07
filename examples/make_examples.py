"""Build the example workbooks.

The numbers are invented and the scenario is fictional. Running this script is the
only way the example files exist, so nothing real can end up in the repository.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).parent


def quarterly_forecast() -> Workbook:
    """A small forecast carrying one of each problem the tool looks for."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Forecast"

    ws["A1"], ws["B1"], ws["C1"], ws["D1"] = "Month", "Units", "Price", "Revenue"
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct"]
    for offset, month in enumerate(months):
        row = offset + 2
        ws.cell(row, 1, month)
        ws.cell(row, 2, 100 + offset * 15)
        ws.cell(row, 3, 42)
        ws.cell(row, 4, f"=B{row}*C{row}")

    # Someone typed last quarter's figure over the formula.
    ws["D5"] = 6800

    # And another cell quietly uses a different price.
    ws["D9"] = "=B9*45"

    # The total was written when the table ended at row 9.
    ws["D12"] = "=SUM(D2:D9)"

    # A rate pulled from a file that may have moved.
    ws["F2"] = "=[1]Assumptions!B3"

    # A stamp that changes every time the file is opened.
    ws["F4"] = "=TODAY()"

    ws.row_dimensions[7].hidden = True

    workings = wb.create_sheet("Workings")
    workings.sheet_state = "hidden"
    workings["A1"] = "Scratch space"

    return wb


def clean_forecast() -> Workbook:
    """The same shape, built properly. The tool should report nothing on this one."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Forecast"
    ws["A1"], ws["B1"], ws["C1"], ws["D1"] = "Month", "Units", "Price", "Revenue"
    for offset in range(10):
        row = offset + 2
        ws.cell(row, 1, f"M{offset + 1}")
        ws.cell(row, 2, 100 + offset * 15)
        ws.cell(row, 3, 42)
        ws.cell(row, 4, f"=B{row}*C{row}")
    ws["D12"] = "=SUM(D2:D11)"
    return wb


if __name__ == "__main__":
    quarterly_forecast().save(HERE / "messy-forecast.xlsx")
    clean_forecast().save(HERE / "clean-forecast.xlsx")
    print(f"wrote messy-forecast.xlsx and clean-forecast.xlsx to {HERE}")
